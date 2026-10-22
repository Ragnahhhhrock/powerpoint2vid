#!/usr/bin/env python3
"""Ingest a .pptx into a job folder the director can work from.

Writes, in <job>/:
  assets/slide-<n>.png     each visible slide rendered by LibreOffice (1280 wide, deck aspect ratio)
  assets/img-<n>-<k>.<ext> pictures embedded in slide n
  plan.json                { title, deck:{width,height,theme}, scenes:[{slide:{title,text,notes,image,images,charts,tables}}] }

Charts are extracted as data (categories + series) so scenes can rebuild them as animated charts with the
exact same numbers. Hidden slides are skipped. Nothing is invented: only what is in the file.

Usage: python3 studio/ingest.py deck.pptx <job_dir> [--max-slides 40] [--title "My deck"]
Needs: python-pptx, LibreOffice (soffice), pdftoppm.
"""
import argparse, json, re, subprocess, sys, tempfile, zipfile
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

MAX_TEXT = 4000          # per slide, keeps requests small
MAX_IMG_BYTES = 6_000_000


def clean(s):
    return re.sub(r"[ \t]+", " ", s.replace("\x0b", "\n").replace("\u2028", "\n")).strip()


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(sh.shapes)
        else:
            yield sh


def theme_colors(path):
    """dk1/lt1/dk2/lt2/accent1-6 hex values from the first theme part, if present."""
    out = {}
    try:
        with zipfile.ZipFile(path) as z:
            name = next((n for n in z.namelist() if re.match(r"ppt/theme/theme\d+\.xml$", n)), None)
            if not name:
                return out
            xml = z.read(name).decode("utf8", "ignore")
        scheme = re.search(r"<a:clrScheme[^>]*>(.*?)</a:clrScheme>", xml, re.S)
        if scheme:
            for tag, body in re.findall(r"<a:(dk1|lt1|dk2|lt2|accent[1-6]|hlink)>(.*?)</a:\1>", scheme.group(1), re.S):
                m = re.search(r'(?:srgbClr val|lastClr)="([0-9A-Fa-f]{6})"', body)
                if m:
                    out[tag] = "#" + m.group(1).lower()
    except Exception:
        pass
    return out


def extract_chart(chart):
    try:
        plot = chart.plots[0]
        cats = [str(c) for c in plot.categories]
        series = [{"name": s.name, "values": [None if v is None else float(v) for v in s.values]} for s in plot.series]
        title = chart.chart_title.text_frame.text if chart.has_title and chart.chart_title.has_text_frame else ""
        return {"type": str(chart.chart_type).split(".")[-1].split(" ")[0].lower(), "title": clean(title), "categories": cats, "series": series}
    except Exception:
        return None


def extract_table(tbl):
    return [[clean(c.text) for c in row.cells] for row in tbl.rows]


def render_slides(deck_path, out_dir, count):
    """LibreOffice -> PDF -> PNGs. Returns the list of PNG paths in slide order."""
    tmp = Path(tempfile.mkdtemp())
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(deck_path)],
                   check=True, capture_output=True, timeout=300)
    pdf = tmp / (Path(deck_path).stem + ".pdf")
    subprocess.run(["pdftoppm", "-png", "-scale-to-x", "1280", "-scale-to-y", "-1", "-l", str(count), str(pdf), str(tmp / "s")],
                   check=True, capture_output=True, timeout=300)
    return sorted(tmp.glob("s-*.png"), key=lambda p: int(p.stem.split("-")[1]))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck"); ap.add_argument("job")
    ap.add_argument("--max-slides", type=int, default=40)
    ap.add_argument("--title", default="")
    a = ap.parse_args()
    job = Path(a.job); (job / "assets").mkdir(parents=True, exist_ok=True)

    prs = Presentation(a.deck)
    W, H = prs.slide_width, prs.slide_height
    slides, hidden = [], 0
    for sl in prs.slides:
        if sl._element.get("show") == "0":
            hidden += 1
            continue
        slides.append(sl)
    if not slides:
        sys.exit("No visible slides in this deck.")
    if len(slides) > a.max_slides:
        sys.exit(f"This deck has {len(slides)} slides; the limit is {a.max_slides}.")

    pngs = render_slides(a.deck, job / "assets", len(slides))
    warn = []
    if len(pngs) != len(slides):
        warn.append(f"rendered {len(pngs)} pages for {len(slides)} slides")

    scenes = []
    for i, sl in enumerate(slides, 1):
        title = clean(sl.shapes.title.text_frame.text) if sl.shapes.title is not None and sl.shapes.title.has_text_frame else ""
        texts, images, charts, tables = [], [], [], []
        for sh in walk(sl.shapes):
            if sh.has_text_frame and sh != sl.shapes.title:
                t = clean(sh.text_frame.text)
                if t:
                    texts.append(t)
            if getattr(sh, "has_chart", False) and sh.has_chart:
                c = extract_chart(sh.chart)
                if c:
                    charts.append(c)
            if getattr(sh, "has_table", False) and sh.has_table:
                tables.append(extract_table(sh.table))
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    blob = sh.image.blob
                    if len(blob) <= MAX_IMG_BYTES:
                        name = f"img-{i}-{len(images) + 1}.{sh.image.ext}"
                        (job / "assets" / name).write_bytes(blob)
                        images.append({"file": name, "w": sh.width, "h": sh.height, "alt": sh._element.xpath(".//p:cNvPr/@descr")[0] if sh._element.xpath(".//p:cNvPr/@descr") else ""})
                except Exception:
                    pass
        notes = clean(sl.notes_slide.notes_text_frame.text) if sl.has_notes_slide and sl.notes_slide.notes_text_frame else ""
        if i - 1 < len(pngs):
            (job / "assets" / f"slide-{i}.png").write_bytes(pngs[i - 1].read_bytes())
        if not title:
            title = (texts[0].split("\n")[0] if texts else f"Slide {i}")[:120]
        scenes.append({"slide": {
            "title": title, "text": "\n".join(texts)[:MAX_TEXT], "notes": notes, "image": f"assets/slide-{i}.png",
            "images": images, "charts": charts, "tables": tables}})

    plan = {"title": a.title or (prs.core_properties.title or Path(a.deck).stem),
            "deck": {"width": int(W), "height": int(H), "theme": theme_colors(a.deck), "hidden_slides_skipped": hidden},
            "scenes": scenes}
    (job / "plan.json").write_text(json.dumps(plan, indent=2))
    print(f"ingested {len(scenes)} slides ({hidden} hidden skipped)" + (f"; warnings: {warn}" if warn else ""))


if __name__ == "__main__":
    main()
