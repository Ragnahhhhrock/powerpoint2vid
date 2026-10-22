#!/usr/bin/env python3
"""Build a deliberately ordinary input deck from <job>/spec.json, render each slide to
<job>/assets/slide-N.png, and write <job>/plan.json (brief + narration per scene).
Usage: python3 studio/examples/build_deck.py <job_dir>
Needs python-pptx, LibreOffice and pdftoppm. Used for the homepage showcase examples."""
import json, subprocess, sys, tempfile
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

job = Path(sys.argv[1]).resolve()
spec = json.loads((job / "spec.json").read_text())
(job / "assets").mkdir(exist_ok=True)

p = Presentation(); p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
for i, sl in enumerate(spec["slides"]):
    chart = sl.get("chart")
    layout = 0 if sl.get("cover") else (5 if chart else 1)
    s = p.slides.add_slide(p.slide_layouts[layout])
    s.shapes.title.text = sl["title"]
    for ph in s.placeholders:   # default layouts are 4:3; stretch placeholders to 16:9
        ph.left, ph.width = int(ph.left * 1.3333), int(ph.width * 1.3333)
    if chart:
        d = CategoryChartData(); d.categories = chart["categories"]; d.add_series(chart["series"], chart["values"])
        kind = XL_CHART_TYPE.LINE if chart.get("type") == "line" else XL_CHART_TYPE.COLUMN_CLUSTERED
        s.shapes.add_chart(kind, Inches(1.5), Inches(1.8), Inches(10.3), Inches(5.2), d)
    elif sl.get("bullets"):
        tf = s.placeholders[1].text_frame; tf.text = sl["bullets"][0]
        for b in sl["bullets"][1:]: tf.add_paragraph().text = b
    s.notes_slide.notes_text_frame.text = sl["notes"]
deck = job / "input-deck.pptx"; p.save(deck)

tmp = Path(tempfile.mkdtemp())
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(deck)], check=True, capture_output=True)
for f in (job / "assets").glob("slide-*.png"): f.unlink()
subprocess.run(["pdftoppm", "-png", "-scale-to-x", "1280", "-scale-to-y", "720", str(tmp / "input-deck.pdf"), str(job / "assets" / "slide")], check=True)
for f in sorted((job / "assets").glob("slide-*.png")):
    f.rename(f.with_name("slide-%d.png" % int(f.stem.split("-")[1])))

plan = {"title": spec["title"], "brief": spec["brief"], "scenes": []}
for i, sl in enumerate(spec["slides"]):
    text = ". ".join(sl.get("bullets", []))
    if sl.get("chart"):
        c = sl["chart"]; text = "; ".join(f"{k}: {v}" for k, v in zip(c["categories"], c["values"])) + f" ({c['series']})"
    sc = {"file": f"scenes/scene-{i}.js", "narration": sl["notes"], "minDuration": sl.get("minDuration", 7),
          "slide": {"title": sl["title"], "text": text, "notes": sl["notes"], "image": f"assets/slide-{i+1}.png"},
          "storyboard": sl.get("storyboard", "")}
    if i and sl.get("transition"): sc["transition"] = sl["transition"]
    plan["scenes"].append(sc)
if "poster" in spec: plan["poster"] = spec["poster"]
(job / "plan.json").write_text(json.dumps(plan, indent=2))
print("ok", job.name, len(spec["slides"]), "slides")
