#!/usr/bin/env python3
"""Build the deliberately ordinary 5-slide input deck for the v3 hero explainer, render its slides
to PNG (assets/slide-1.png ...), and write plan.json (brief + narration + storyboard per scene).

The deck is the plain "before": default Office theme, default layouts, one default chart, and wordy
speaker notes the way people actually write them. The plan's narration is the tightened version
(written by a.i. in a Claude session from these notes), and scene 1 shows that tightening on screen.
Run from this folder. Needs python-pptx, LibreOffice, pdftoppm."""
import json, shutil, subprocess, tempfile
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

HERE = Path(__file__).parent
(HERE / "assets").mkdir(exist_ok=True)

# (title, bullets, chart categories or None, speaker notes as written in the deck)
SLIDES = [
    ("PowerPoint2Vid", ["Turn slides into explainer videos"], None,
     "OK so this first one is just a basic title slide, pretty standard. What I want to show is "
     "what it looks like once it gets properly directed."),
    ("How it works", ["Reads the whole deck", "Tightens the story", "Speaker notes become the voiceover",
                      "Each slide becomes an animated scene"], None,
     "So basically how it works is the tool reads through all of the slides in the deck, and then it "
     "goes through and tightens up the story so it flows. The speaker notes that you have written get "
     "used as the voiceover. Then each slide gets redesigned as an animated scene, with the animation "
     "timed to the narration."),
    ("Quarterly revenue", [], ("Q1", "Q2", "Q3", "Q4"),
     "Here is the revenue chart for the four quarters. As you can see it goes up every quarter, from "
     "1.2 million in Q1 to 2.6 million in Q4."),
    ("Plans", ["Basic: exports carry a PowerPoint2Vid watermark", "Pro: from A$29/month",
               "Pro includes 10 a.i.-directed videos, no watermark"], None,
     "In terms of plans, Basic exports have a watermark on them. Pro is from A$29 a month and includes "
     "10 a.i.-directed videos with no watermark."),
    ("Get started", ["powerpoint2vid.com"], None,
     "To get started, head over to powerpoint2vid.com."),
]
REVENUE = (1.2, 1.5, 1.9, 2.6)

p = Presentation(); p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
for i, (title, bullets, chart, notes) in enumerate(SLIDES):
    s = p.slides.add_slide(p.slide_layouts[0 if i in (0, 4) else (5 if chart else 1)])
    s.shapes.title.text = title
    for ph in s.placeholders:   # default layouts are 4:3; stretch placeholders to 16:9
        l, t, w, h = ph.left, ph.top, ph.width, ph.height
        ph.left, ph.top, ph.width, ph.height = int(l * 1.3333), t, int(w * 1.3333), h
    if chart:
        d = CategoryChartData(); d.categories = chart; d.add_series("Revenue ($M)", REVENUE)
        s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1.5), Inches(1.8), Inches(10.3), Inches(5.2), d)
    elif bullets:
        tf = s.placeholders[1].text_frame; tf.text = bullets[0]
        for b in bullets[1:]: tf.add_paragraph().text = b
    s.notes_slide.notes_text_frame.text = notes
deck = HERE / "input-deck.pptx"; p.save(deck)

tmp = Path(tempfile.mkdtemp())
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(deck)], check=True, capture_output=True)
for f in (HERE / "assets").glob("slide-*.png"): f.unlink()
subprocess.run(["pdftoppm", "-png", "-scale-to-x", "1280", "-scale-to-y", "720", str(tmp / "input-deck.pdf"), str(HERE / "assets" / "slide")], check=True)
for f in sorted((HERE / "assets").glob("slide-*.png")):   # slide-1.png .. (pdftoppm may zero-pad)
    f.rename(f.with_name("slide-%d.png" % int(f.stem.split("-")[1])))
icon = HERE.parents[2] / "assets" / "icon-512.png"
if icon.exists(): shutil.copy(icon, HERE / "assets" / "icon.png")

# The video: 6 scenes from 5 slides (slide 2 is split into "story" and "redesign").
SCENES = [
    dict(slide=0, minDuration=7.6,
         narration="This is an ordinary slide. Watch it get directed.",
         storyboard="The real default title slide fills the frame in a plain editor window. On 'directed' a "
                    "clapper-snap flash, then the slide breaks into a wave of tiles that tumble toward camera, "
                    "revealing a dark stage with a projector beam and dust. The title is rebuilt as heavy kinetic "
                    "type, letters rising through a mask, an amber light sweep crossing them."),
    dict(slide=1, minDuration=9.4, transition={"type": "fadeblack", "duration": 0.45},
         narration="a.i. reads your whole deck and tightens the story. Your speaker notes become the voiceover.",
         storyboard="The five real slides fan in a 3D arc; an amber scan line reads across them on 'reads'. The deck's "
                    "actual wordy speaker note is typeset below; on 'tightens' filler words are struck in ember and "
                    "collapse out, the remaining words close up. On 'voiceover' the tightened line lifts and a "
                    "waveform grows from it, a playhead sweeping across."),
    dict(slide=1, minDuration=9.0, transition={"type": "smoothleft", "duration": 0.6},
         narration="Then it redesigns every slide as an animated scene, with each move timed to the narration.",
         storyboard="The plain bullet slide, as written. On 'redesigns' the camera swings to an isometric exploded "
                    "view: the new scene's layers (stage, light, type, cards) float apart in depth with mono labels. "
                    "On 'animated scene' the layers slam together and the camera returns front-on. On 'timed' a "
                    "timeline strip draws along the bottom: waveform, keyframe diamonds, a playhead that pops each "
                    "step card as it passes."),
    dict(slide=2, minDuration=9.2, transition={"type": "fadeblack", "duration": 0.45},
         narration="A flat chart becomes a moment. The numbers build, and the result lands as it's said.",
         storyboard="The deck's default chart, as-is, tips back on 'flat' and becomes the floor. On 'moment' the same "
                    "data erupts from it as extruded 3D columns under an orbiting camera. Counters build column by "
                    "column on 'numbers build'; on 'lands' Q4 ignites ember, the rest dim, and '+117%' (Q4 vs Q1, "
                    "from the deck's numbers) slams in."),
    dict(slide=3, minDuration=10.6, transition={"type": "smoothleft", "duration": 0.6},
         narration="Basic exports carry a watermark. Pro removes it, and includes ten a.i.-directed videos a month, from 29 Australian dollars.",
         storyboard="A real frame from this video in a player, stamped with the PowerPoint2Vid watermark, labelled "
                    "BASIC. On 'Pro removes it' a light wipe clears the watermark and the frame becomes the hero, "
                    "labelled PRO. On 'ten' ten frames from this video deal out into a film strip while a counter "
                    "climbs to 10. On 'Australian dollars' the price A$29/month lands in big type."),
    dict(slide=4, minDuration=8.6, transition={"type": "fadewhite", "duration": 0.6},
         narration="More than a PowerPoint to video converter. Go Pro at powerpoint2vid.com.",
         storyboard="Finale under the projector beam. 'More than a PowerPoint to video converter.' types in muted; "
                    "then 'Your deck, directed.' rises huge with an amber light sweep. On 'Go Pro' the Pro button "
                    "lands with the price, and the URL settles beneath. Holds as an end card."),
]

plan = {
    "title": "PowerPoint2Vid: your deck, directed",
    "poster": 7.0,   # seconds into the final video: the settled title card in scene 0
    "brief": {
        "font": "Inter",
        "palette": {"ink": "#0c0d10", "night": "#060608", "coal": "#121318", "paper": "#f4eee3",
                    "amber": "#ffb340", "ember": "#ff5a2c", "teal": "#3fd9c6", "muted": "#8c8e98"},
        "motion": "Cinematic and decisive: expo.inOut camera moves, back.out pops with small overshoot, a slow "
                  "continuous drift so no frame is static, light sweeps and dust in a projector beam. Every key "
                  "visual lands on the word that names it.",
        "story": "Before and after, escalating: a plain white slide is directed into a dark film set (teal and "
                 "amber grade, like a film), then each scene adds a new craft: story, layered redesign, 3D data, "
                 "and the Pro offer, ending under the same beam it opened with."
    },
    "scenes": []
}
for i, d in enumerate(SCENES):
    title, bullets, chart, notes = SLIDES[d["slide"]]
    sc = {"file": f"scenes/scene-{i}.js", "slide_index": d["slide"], "narration": d["narration"],
          "minDuration": d["minDuration"], "storyboard": d["storyboard"],
          "slide": {"title": title, "text": ". ".join(bullets), "notes": notes, "image": f"assets/slide-{d['slide'] + 1}.png"}}
    if d.get("transition"): sc["transition"] = d["transition"]
    plan["scenes"].append(sc)
(HERE / "plan.json").write_text(json.dumps(plan, indent=2))
print("ok")
