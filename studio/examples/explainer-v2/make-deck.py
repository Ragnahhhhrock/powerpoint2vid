#!/usr/bin/env python3
"""Build the deliberately ordinary 5-slide input deck for the v2 explainer, render its slides
to PNG (assets/slide-1.png ...), and write plan.json (brief + narration per scene).
Run from this folder. Needs python-pptx, LibreOffice, pdftoppm."""
import json, subprocess, tempfile
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

HERE = Path(__file__).parent
SLIDES = [
    ("PowerPoint2Vid", ["Turn slides into explainer videos"], None,
     "This video started life as five ordinary slides. Watch what happens when they get a director."),
    ("The problem", ["Most slide videos are just slides and a voice", "The bullet points get read aloud", "Audiences tune out fast"], None,
     "Most slide videos are just slides, with a voice reading out the bullet points. By slide three, your audience has tuned out."),
    ("How it works", ["Reads your whole deck", "Writes the story", "Designs every scene", "Times animation to the narration"], None,
     "PowerPoint2Vid reads your whole deck. It writes the story, designs every scene, and times each animation to the narration."),
    ("Quarterly revenue", [], ("Q1", "Q2", "Q3", "Q4"),
     "So a flat chart becomes a moment. The numbers build, and the result lands right as it's said."),
    ("Get started", ["powerpoint2vid.com"], None,
     "Turn ordinary slides into extraordinary explainer videos. Start at powerpoint2vid.com."),
]

p = Presentation(); p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
for i, (title, bullets, chart, notes) in enumerate(SLIDES):
    s = p.slides.add_slide(p.slide_layouts[0 if i in (0, 4) else (5 if chart else 1)])
    s.shapes.title.text = title
    for ph in s.placeholders:   # default layouts are 4:3; stretch placeholders to 16:9
        l, t, w, h = ph.left, ph.top, ph.width, ph.height
        ph.left, ph.top, ph.width, ph.height = int(l * 1.3333), t, int(w * 1.3333), h
    if chart:
        d = CategoryChartData(); d.categories = chart; d.add_series("Revenue ($M)", (1.2, 1.5, 1.9, 2.6))
        s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1.5), Inches(1.8), Inches(10.3), Inches(5.2), d)
    elif bullets:
        tf = s.placeholders[1].text_frame; tf.text = bullets[0]
        for b in bullets[1:]: tf.add_paragraph().text = b
    s.notes_slide.notes_text_frame.text = notes
deck = HERE / "input-deck.pptx"; p.save(deck)

tmp = Path(tempfile.mkdtemp())
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(deck)], check=True, capture_output=True)
subprocess.run(["pdftoppm", "-png", "-scale-to-x", "1280", "-scale-to-y", "720", str(tmp / "input-deck.pdf"), str(HERE / "assets" / "slide")], check=True)
for f in sorted((HERE / "assets").glob("slide-*.png")):   # slide-1.png .. (pdftoppm may zero-pad)
    f.rename(f.with_name("slide-%d.png" % int(f.stem.split("-")[1])))

plan = {
    "title": "PowerPoint2Vid: ordinary slides, extraordinary videos",
    "brief": {
        "font": "Inter",
        "palette": {"ink": "#0e1330", "night": "#080b1f", "violet": "#6d4aff", "violet-hi": "#8f75ff",
                    "coral": "#ff6b4a", "mint": "#1fbf8f", "paper": "#f8f7fc", "muted": "#aab0d4"},
        "motion": "Confident and cinematic: slow camera push, springy reveals with a small overshoot, "
                  "light streaks on key beats, film grain. Every key visual lands on the word that names it.",
        "story": "Before/after: plain office slides are shown as they are, then transformed. Dark stage for "
                 "drama, one light paper scene for the how-it-works explanation, gradient finale."
    },
    "scenes": []
}
trans = [None, {"type": "fadeblack", "duration": 0.5}, {"type": "circleopen", "duration": 0.7},
         {"type": "smoothleft", "duration": 0.6}, {"type": "fade", "duration": 0.6}]
for i, (title, bullets, chart, notes) in enumerate(SLIDES):
    sc = {"file": f"scenes/scene-{i}.js", "narration": notes, "minDuration": 6,
          "slide": {"title": title, "text": ". ".join(bullets), "notes": notes, "image": f"assets/slide-{i+1}.png"}}
    if trans[i]: sc["transition"] = trans[i]
    plan["scenes"].append(sc)
(HERE / "plan.json").write_text(json.dumps(plan, indent=2))
print("ok")
